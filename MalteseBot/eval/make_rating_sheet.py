"""make_rating_sheet.py — build the BLIND human-rating workbook for judge validation.

Samples answers from results_multimodel.csv (run 1): every answer of the primary
model (GPT-4o, both conditions) plus a stratified random sample of the other
models, balanced over condition and language. Rows are shuffled and carry no
model or condition label; the mapping is written separately to rating_key.json.

The rater marks, for each answer:
    incorrect_claim  1 if the answer states anything about the law that is wrong,
                     fabricated or contradicts the reference / key facts; else 0
    omission         1 if one or more key facts are missing; else 0
    note             optional free text

Run (from MalteseBot, inside the venv):
    .venv\\Scripts\\python.exe eval\\make_rating_sheet.py --extra 72
Outputs: eval/rating_sheet.xlsx (give to the rater), eval/rating_key.json (keep back)
"""
from __future__ import annotations

import argparse
import csv
import json
import random
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EVAL_DIR = Path(__file__).resolve().parent
RESULTS = EVAL_DIR / "results_multimodel.csv"
GOLD = EVAL_DIR / "gold_queries.json"
PRIMARY = "gpt-4o-2024-11-20"
SEED = 20260906

INSTRUCTIONS = [
    "BLIND RATING SHEET — MalteseLegalBot technical evaluation (resubmission, September 2026)",
    "",
    "Each row is one answer produced by one of six language models, with or without retrieval. "
    "You do not know which; do not try to guess. Rate the ANSWER against the REFERENCE and KEY FACTS, "
    "which were verified against the consolidated statutes on legislation.mt.",
    "",
    "Column incorrect_claim: enter 1 if the answer states ANYTHING about the law that is wrong, invented, "
    "or contradicts the reference or key facts (a wrong number, a wrong rule, an invented article, penalty "
    "or procedure). Correct extra detail does not count. Disclaimers do not count. Leaving something out "
    "does NOT count here. Otherwise enter 0.",
    "",
    "Column omission: enter 1 if one or more of the key facts is missing from the answer. Otherwise 0.",
    "",
    "Column note: optional. Say which statement is wrong or which fact is missing.",
    "",
    "Work in order, one row at a time, without looking back at earlier rows. Do not change any other column. "
    "Save the file with the same name when finished.",
]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--extra", type=int, default=72, help="answers sampled from the non-primary models")
    ap.add_argument("--run", default="1")
    args = ap.parse_args()

    gold = {q["id"]: q for q in json.loads(GOLD.read_text(encoding="utf-8"))["queries"]}
    with open(RESULTS, encoding="utf-8") as f:
        rows = [r for r in csv.DictReader(f) if r["run"] == args.run]
    primary = [r for r in rows if r["model"] == PRIMARY]
    others = [r for r in rows if r["model"] != PRIMARY]

    # Stratify the extra sample over (model, condition, lang) as evenly as possible.
    rng = random.Random(SEED)
    strata: dict[tuple, list] = defaultdict(list)
    for r in others:
        strata[(r["model"], r["condition"], r["lang"])].append(r)
    keys = sorted(strata)
    for k in keys:
        rng.shuffle(strata[k])
    sample: list[dict] = []
    while len(sample) < args.extra and any(strata[k] for k in keys):
        for k in keys:
            if strata[k] and len(sample) < args.extra:
                sample.append(strata[k].pop())

    selected = primary + sample
    rng.shuffle(selected)

    key = {}
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        c = ws0.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws0.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Ratings")
    header = ["rating_id", "question", "reference_answer", "key_facts", "answer",
              "incorrect_claim", "omission", "note"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    widths = [10, 40, 45, 25, 70, 14, 10, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for n, r in enumerate(selected, start=1):
        rid = f"R{n:03d}"
        q = gold[r["id"]]
        ws.append([rid, q["question"], q["reference_answer"], "; ".join(q["key_facts"]),
                   r["answer"], None, None, None])
        for c in ws[n + 1]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        key[rid] = {"model": r["model"], "condition": r["condition"], "run": r["run"], "id": r["id"],
                    "halluc_j1": r["halluc_j1"], "omission_j1": r["omission_j1"],
                    "halluc_j2": r["halluc_j2"], "omission_j2": r["omission_j2"]}
    ws.freeze_panes = "B2"

    out = EVAL_DIR / "rating_sheet.xlsx"
    wb.save(out)
    (EVAL_DIR / "rating_key.json").write_text(json.dumps(key, indent=2), encoding="utf-8")
    by_model = defaultdict(int)
    for v in key.values():
        by_model[v["model"]] += 1
    print(f"wrote {out.name} with {len(selected)} answers; key -> rating_key.json")
    print("per model:", dict(by_model))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
