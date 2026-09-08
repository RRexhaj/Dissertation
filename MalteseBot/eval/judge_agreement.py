"""judge_agreement.py -- validation of the two LLM judges against a third,
cross-vendor judge and human adjudication of every disagreement.

Inputs
    rating_key.json        the blind sample (120 answers) with the verdicts of
                           judge 1 (gpt-4.1) and judge 2 (gpt-4o) and the
                           model/condition of every row
    rating_claude.csv      verdicts of judge 3 (Claude, Anthropic), given blind
                           on rating_sheet.xlsx with the same rubric
    adjudication_sheet.xlsx  (optional) the rows on which the three judges were
                           not unanimous, with the author's final verdict in
                           the incorrect_claim / omission columns

The reference verdict for a row is the unanimous verdict where all three judges
agree and the adjudicated verdict otherwise.  Rows still awaiting adjudication
are excluded from the accuracy figures (and counted).

Run (from MalteseBot, inside the venv):
    .venv\\Scripts\\python.exe eval\\judge_agreement.py --make-adjudication
    ... the author fills eval/adjudication_sheet.xlsx ...
    .venv\\Scripts\\python.exe eval\\judge_agreement.py
Outputs: eval/judge_agreement.json, eval/tables/tab_judge_agreement.tex,
         eval/tables/tab_judge_confusion.tex, eval/adjudication_sheet.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EVAL_DIR = Path(__file__).resolve().parent
TABLES = EVAL_DIR / "tables"
SHEET = EVAL_DIR / "rating_sheet.xlsx"
KEY = EVAL_DIR / "rating_key.json"
THIRD = EVAL_DIR / "rating_claude.csv"
ADJ = EVAL_DIR / "adjudication_sheet.xlsx"

JUDGES = ("j1", "j2", "j3")
JUDGE_NAMES = {"j1": "Judge 1 (GPT-4.1)", "j2": "Judge 2 (GPT-4o)", "j3": "Judge 3 (Claude)"}
METRICS = (("halluc", "Incorrect claim"), ("omission", "Omission"))

ADJ_INSTRUCTIONS = [
    "ADJUDICATION SHEET -- judge validation (resubmission, September 2026)",
    "",
    "Each row is an answer on which the three automated judges did NOT agree on at least one of the two "
    "flags. Their verdicts are hidden. Read the ANSWER against the REFERENCE and KEY FACTS (verified "
    "against legislation.mt) and give your own final verdict.",
    "",
    "incorrect_claim: 1 if the answer states anything about the law that is wrong, invented, or contradicts "
    "the reference or key facts; correct extra detail and disclaimers do not count; leaving something out "
    "does not count here. Otherwise 0.",
    "omission: 1 if one or more key facts is missing (an equivalent formulation, e.g. 0.2 g/l for 20 mg/100 ml, "
    "counts as present). Otherwise 0.",
    "note: optional.",
    "",
    "Fill BOTH columns for every row even if only one flag was disputed. Save with the same name.",
]


def cohen_kappa(a: list[int], b: list[int]) -> float:
    n = len(a)
    if n == 0:
        return float("nan")
    po = sum(1 for x, y in zip(a, b) if x == y) / n
    pa1, pb1 = sum(a) / n, sum(b) / n
    pe = pa1 * pb1 + (1 - pa1) * (1 - pb1)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def fleiss_kappa(rows: list[list[int]]) -> float:
    """rows: one list of binary verdicts per item (all items rated by the same k raters)."""
    n = len(rows)
    if n == 0:
        return float("nan")
    k = len(rows[0])
    p_i = []
    for r in rows:
        c1 = sum(r)
        c0 = k - c1
        p_i.append((c1 * c1 + c0 * c0 - k) / (k * (k - 1)))
    p_bar = sum(p_i) / n
    p1 = sum(sum(r) for r in rows) / (n * k)
    pe = p1 * p1 + (1 - p1) * (1 - p1)
    return 1.0 if pe == 1 else (p_bar - pe) / (1 - pe)


def confusion(truth: list[int], pred: list[int]) -> dict:
    tp = sum(1 for t, p in zip(truth, pred) if t == 1 and p == 1)
    tn = sum(1 for t, p in zip(truth, pred) if t == 0 and p == 0)
    fp = sum(1 for t, p in zip(truth, pred) if t == 0 and p == 1)
    fn = sum(1 for t, p in zip(truth, pred) if t == 1 and p == 0)
    sens = tp / (tp + fn) if tp + fn else None
    spec = tn / (tn + fp) if tn + fp else None
    acc = (tp + tn) / len(truth) if truth else None
    return {"tp": tp, "tn": tn, "fp": fp, "fn": fn,
            "sensitivity": None if sens is None else round(sens, 3),
            "specificity": None if spec is None else round(spec, 3),
            "accuracy": None if acc is None else round(acc, 3)}


def landis_koch(k: float) -> str:
    if k != k:
        return "n/a"
    for lim, lab in [(0.0, "poor"), (0.2, "slight"), (0.4, "fair"), (0.6, "moderate"), (0.8, "substantial")]:
        if k <= lim:
            return lab
    return "almost perfect"


def load_verdicts() -> tuple[dict, dict]:
    key = json.loads(KEY.read_text(encoding="utf-8"))
    third: dict[str, dict] = {}
    with open(THIRD, encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            third[r["rating_id"]] = {"halluc": int(r["incorrect_claim"]), "omission": int(r["omission"]),
                                     "note": r.get("note", "")}
    verdicts: dict[str, dict] = {}
    for rid, k in key.items():
        if rid not in third:
            continue
        verdicts[rid] = {
            "condition": k["condition"], "model": k["model"], "id": k["id"],
            "halluc": {"j1": int(k["halluc_j1"]), "j2": int(k["halluc_j2"]), "j3": third[rid]["halluc"]},
            "omission": {"j1": int(k["omission_j1"]), "j2": int(k["omission_j2"]), "j3": third[rid]["omission"]},
            "note_j3": third[rid]["note"],
        }
    return key, verdicts


def load_adjudication() -> dict[str, dict]:
    if not ADJ.exists():
        return {}
    ws = load_workbook(ADJ, data_only=True)["Adjudication"]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header)}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rid = row[col["rating_id"]]
        ic, om = row[col["incorrect_claim"]], row[col["omission"]]
        if rid is None or ic is None or om is None:
            continue
        out[rid] = {"halluc": int(ic), "omission": int(om)}
    return out


def disputed(verdicts: dict) -> list[str]:
    return [rid for rid, v in verdicts.items()
            if len(set(v["halluc"].values())) > 1 or len(set(v["omission"].values())) > 1]


def make_adjudication_sheet(verdicts: dict) -> int:
    rows = disputed(verdicts)
    sheet = load_workbook(SHEET, data_only=True)["Ratings"]
    header = [c.value for c in sheet[1]]
    col = {h: i for i, h in enumerate(header)}
    text = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        text[row[col["rating_id"]]] = row
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instructions"
    for i, line in enumerate(ADJ_INSTRUCTIONS, start=1):
        c = ws0.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws0.column_dimensions["A"].width = 120
    ws = wb.create_sheet("Adjudication")
    ws.append(["rating_id", "question", "reference_answer", "key_facts", "answer",
               "incorrect_claim", "omission", "note"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for i, w in enumerate([10, 40, 45, 25, 70, 14, 10, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for n, rid in enumerate(rows, start=2):
        r = text[rid]
        ws.append([rid, r[col["question"]], r[col["reference_answer"]], r[col["key_facts"]],
                   r[col["answer"]], None, None, None])
        for c in ws[n]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    wb.save(ADJ)
    return len(rows)


def tex_escape(s: str) -> str:
    return s.replace("&", "\\&").replace("%", "\\%")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--make-adjudication", action="store_true",
                    help="write eval/adjudication_sheet.xlsx with the non-unanimous rows (verdicts hidden)")
    args = ap.parse_args()

    key, verdicts = load_verdicts()
    rids = sorted(verdicts)
    print(f"{len(rids)} of {len(key)} answers rated by all three judges")
    if args.make_adjudication:
        n = make_adjudication_sheet(verdicts)
        print(f"adjudication sheet written with {n} disputed rows -> {ADJ.name}")

    adj = load_adjudication()
    disp = disputed(verdicts)
    out: dict = {"n_rated": len(rids), "n_total": len(key), "n_disputed": len(disp),
                 "n_adjudicated": sum(1 for r in disp if r in adj), "pairs": {}, "judges": {}, "fleiss": {},
                 "disputed_by_condition": dict(Counter(verdicts[r]["condition"] for r in disp))}

    for metric, _label in METRICS:
        v = {j: [verdicts[r][metric][j] for r in rids] for j in JUDGES}
        pairs = {}
        for a, b in (("j1", "j2"), ("j1", "j3"), ("j2", "j3")):
            k = cohen_kappa(v[a], v[b])
            pairs[f"{a}_vs_{b}"] = {"kappa": round(k, 3), "interpretation": landis_koch(k), "n": len(rids),
                                    "agreement": round(sum(1 for x, y in zip(v[a], v[b]) if x == y) / len(rids), 3)}
        out["pairs"][metric] = pairs
        fk = fleiss_kappa([[verdicts[r][metric][j] for j in JUDGES] for r in rids])
        out["fleiss"][metric] = {"kappa": round(fk, 3), "interpretation": landis_koch(fk)}
        out[f"{metric}_disputed"] = sum(1 for r in rids if len(set(verdicts[r][metric].values())) > 1)
        for j in JUDGES:
            out[f"{metric}_rate_{j}"] = round(sum(v[j]) / len(rids), 3)

        # Reference verdict: unanimous, else adjudicated (rows awaiting adjudication are dropped).
        ref_ids, ref = [], []
        for r in rids:
            vals = set(verdicts[r][metric].values())
            if len(vals) == 1:
                ref_ids.append(r)
                ref.append(vals.pop())
            elif r in adj:
                ref_ids.append(r)
                ref.append(adj[r][metric])
        out[f"{metric}_n_reference"] = len(ref_ids)
        out[f"{metric}_rate_reference"] = round(sum(ref) / len(ref), 3) if ref else None
        for cond in ("RAG", "Baseline"):
            sub = [x for r, x in zip(ref_ids, ref) if verdicts[r]["condition"] == cond]
            out[f"{metric}_rate_reference_{cond}"] = round(sum(sub) / len(sub), 3) if sub else None
        judges = {}
        for j in JUDGES:
            pred = [verdicts[r][metric][j] for r in ref_ids]
            judges[j] = confusion(ref, pred)
            judges[j]["kappa_vs_reference"] = round(cohen_kappa(ref, pred), 3)
        out["judges"][metric] = judges

    # Where the adjudicator sided, per judge (how often each judge was right on disputed rows)
    if adj:
        sided = {}
        for metric, _ in METRICS:
            sided[metric] = {}
            for j in JUDGES:
                rows = [r for r in disp if r in adj and len(set(verdicts[r][metric].values())) > 1]
                if rows:
                    sided[metric][j] = {"n": len(rows),
                                        "agreed_with_adjudicator": sum(1 for r in rows if verdicts[r][metric][j] == adj[r][metric])}
        out["disputed_sided"] = sided

    (EVAL_DIR / "judge_agreement.json").write_text(json.dumps(out, indent=2), encoding="utf-8")

    # ---- LaTeX tables -------------------------------------------------------------------------
    TABLES.mkdir(exist_ok=True)
    n_ref = out["halluc_n_reference"]
    adj_note = ("the adjudicated verdict where they disagreed" if adj else
                "PENDING ADJUDICATION: disputed rows are excluded")
    lines = [
        "\\begin{table}[H]",
        "\\centering",
        "{\\footnotesize\\setstretch{1.0}\\setlength{\\tabcolsep}{4pt}\\def\\arraystretch{1.2}",
        "\\caption{Judge validation on the blind sample of " + str(len(rids)) + " answers. Pairwise Cohen's "
        "$\\kappa$ between the two GPT judges and the cross-vendor third judge (Claude); sensitivity, specificity "
        "and accuracy of each judge against the reference verdict, which is the unanimous verdict where the three "
        "judges agreed and " + adj_note + " (n = " + str(n_ref) + ").}",
        "\\label{tab:judge-agreement}",
        "\\begin{tabular}{llcccc}",
        "\\toprule",
        "Criterion & Comparison & Cohen's $\\kappa$ & Agreement & Sensitivity & Specificity \\\\",
        "\\midrule",
    ]
    pair_names = {"j1_vs_j2": "Judge 1 vs judge 2", "j1_vs_j3": "Judge 1 vs judge 3", "j2_vs_j3": "Judge 2 vs judge 3"}
    for metric, label in METRICS:
        first = True
        for name, p in out["pairs"][metric].items():
            lines.append(f"{label if first else ''} & {pair_names[name]} & {p['kappa']:.2f} ({p['interpretation']}) & "
                         f"{p['agreement']:.2f} & -- & -- \\\\")
            first = False
        fk = out["fleiss"][metric]
        lines.append(f" & Three judges (Fleiss) & {fk['kappa']:.2f} ({fk['interpretation']}) & -- & -- & -- \\\\")
        for j in JUDGES:
            c = out["judges"][metric][j]
            se = "--" if c["sensitivity"] is None else f"{c['sensitivity']:.2f}"
            sp = "--" if c["specificity"] is None else f"{c['specificity']:.2f}"
            lines.append(f" & {JUDGE_NAMES[j]} vs reference & {c['kappa_vs_reference']:.2f} & {c['accuracy']:.2f} & {se} & {sp} \\\\")
        if metric == "halluc":
            lines.append("\\midrule")
    lines += ["\\bottomrule", "\\end{tabular}}", "\\end{table}"]
    (TABLES / "tab_judge_agreement.tex").write_text("\n".join(lines) + "\n", encoding="utf-8")

    # Confusion counts for Appendix F
    lines = [
        "\\begin{table}[H]",
        "\\centering",
        "{\\footnotesize\\setstretch{1.0}\\setlength{\\tabcolsep}{4pt}\\def\\arraystretch{1.2}",
        "\\caption{Confusion counts of each judge against the reference verdict (n = " + str(n_ref) + "; "
        "TP = flagged and reference flagged, FP = flagged but reference clear, FN = missed, TN = both clear).}",
        "\\label{tab:judge-confusion}",
        "\\begin{tabular}{llcccc}",
        "\\toprule",
        "Criterion & Judge & TP & FP & FN & TN \\\\",
        "\\midrule",
    ]
    for metric, label in METRICS:
        for i, j in enumerate(JUDGES):
            c = out["judges"][metric][j]
            lines.append(f"{label if i == 0 else ''} & {JUDGE_NAMES[j]} & {c['tp']} & {c['fp']} & {c['fn']} & {c['tn']} \\\\")
        if metric == "halluc":
            lines.append("\\midrule")
    lines += ["\\bottomrule", "\\end{tabular}}", "\\end{table}"]
    (TABLES / "tab_judge_confusion.tex").write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(json.dumps({k: v for k, v in out.items() if k not in ("judges",)}, indent=2))
    print("judges:", json.dumps(out["judges"], indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
